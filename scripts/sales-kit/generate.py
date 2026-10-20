from pathlib import Path
import argparse, base64, html, json, subprocess, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Genera el kit comercial de Nexo.')
parser.add_argument('--text-only',action='store_true',help='Actualiza textos, Word y ZIP conservando el PDF y el JPG existentes.')
args=parser.parse_args()
KIT=ROOT/'public/SKYROTE13/kit-comercial'
PRESENT=ROOT/'public/presentacion'
CONTENT=json.loads((Path(__file__).parent/'content.json').read_text())
KIT.mkdir(parents=True,exist_ok=True);PRESENT.mkdir(parents=True,exist_ok=True)
if args.text_only:
 for filename in ['Nexo-Informacion-Conectada.pdf','Nexo-WhatsApp-Informacion.jpg']:
  if not (PRESENT/filename).is_file():
   parser.error(f'--text-only necesita el adjunto existente: {filename}')
KIT.joinpath('Investigacion-copy-Nexo.md').write_bytes((ROOT/'docs/comercial/investigacion-copy.md').read_bytes())
e=html.escape
logo='data:image/svg+xml;base64,'+base64.b64encode((ROOT/'public/brand/nexo-logo.svg').read_bytes()).decode()
font=base64.b64encode((ROOT/'src/fonts/dm-sans-latin.woff2').read_bytes()).decode()
heading=base64.b64encode((ROOT/'src/fonts/manrope-latin.woff2').read_bytes()).decode()
PRESENT.joinpath('nexo-alquileres.html').write_text(f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow"><title>Nexo | Equipos y alquileres</title><link rel="icon" href="/favicon.ico"><style>
@font-face{{font-family:DM;src:url(data:font/woff2;base64,{font})}}@font-face{{font-family:Manrope;src:url(data:font/woff2;base64,{heading})}}*{{box-sizing:border-box}}@page{{size:A4;margin:0}}body{{margin:0;background:#edf2ee;color:#202730;font:11pt/1.65 DM,sans-serif}}.sheet{{width:210mm;min-height:297mm;margin:auto;background:white;padding:14mm 19mm 12mm}}header{{display:flex;align-items:center;justify-content:space-between;padding-bottom:7mm;border-bottom:1px solid #dce4df}}header img{{width:33mm}}header span{{font-size:9pt;color:#607168;text-align:right}}h1,h2{{font-family:Manrope,sans-serif;letter-spacing:-.035em}}h1{{font-size:28pt;line-height:1.18;margin:7mm 0 5mm;max-width:160mm;font-weight:700}}h1 span{{color:#247459}}.intro{{font-size:11.5pt;color:#52675b;max-width:163mm;margin:0 0 6mm}}h2{{font-size:14pt;line-height:1.4;margin:0 0 2mm}}p{{margin:0}}.scenario{{padding:6mm;background:#eef2ee;border-radius:2mm}}.scenario>p{{font-size:10pt;color:#52675b}}.flow{{display:grid;grid-template-columns:1fr 1fr 1fr;gap:4mm;margin-top:5mm}}.flow>div{{background:white;padding:4mm;border-radius:1.5mm}}.flow strong{{display:block;color:#247459;font-size:10.5pt;margin-bottom:1mm}}.flow span{{display:block;font-size:9pt;color:#52675b;line-height:1.55}}.details{{padding:5mm 0 0}}.details ul{{padding:0;list-style:none;margin:3mm 0 0}}.details li{{display:grid;grid-template-columns:46mm 1fr;gap:5mm;padding:2mm 0;border-bottom:1px solid #e1e7e3;font-size:10pt}}.details li strong{{font-weight:600}}.details li span{{color:#52675b}}.meeting{{margin-top:5mm;background:#20352b;color:white;padding:5mm 6mm;border-radius:2mm}}.meeting h2{{color:white;font-size:15pt;margin:0 0 2mm}}.meeting p{{font-size:10pt;line-height:1.65;color:#e0ece4}}.meeting a{{color:white}}.scope{{font-size:8.5pt;color:#607168;line-height:1.65;margin-top:4mm}}footer{{border-top:1px solid #dce4df;padding-top:4mm;margin-top:3mm;display:flex;justify-content:space-between;gap:5mm;font-size:9pt;line-height:1.8}}footer a{{color:#247459;text-decoration:none}}footer span{{color:#607168}}@media screen and (max-width:800px){{.sheet{{width:100%;min-height:0;padding:24px}}h1{{font-size:32px}}.flow{{grid-template-columns:1fr}}.details li{{grid-template-columns:1fr;gap:2px}}footer{{flex-wrap:wrap}}}}@media print{{body{{background:white}}.sheet{{margin:0}}}}
</style></head><body><main class="sheet"><header><img src="{logo}" alt="Nexo"><span>Sistemas de información a medida<br>Empresas de alquiler de equipos</span></header><h1>La información de cada equipo.<br><span>Conectada con su operación.</span></h1><p class="intro">Relaciona clientes, equipos, contratos y mantenimiento en un sistema diseñado alrededor de cómo trabaja tu empresa.</p><section class="scenario"><h2>Una ficha para consultar la historia completa.</h2><p>Un ejemplo aplicado al alquiler de impresoras, herramientas y equipos audiovisuales.</p><div class="flow"><div><strong>Equipo</strong><span>Identificación, estado y cliente relacionado.</span></div><div><strong>Documentos</strong><span>Contrato y reportes asociados al equipo.</span></div><div><strong>Mantenimiento</strong><span>Historial de servicios y sus documentos PDF.</span></div></div></section><section class="details"><h2>¿Qué podemos revisar juntos?</h2><ul><li><strong>Tu proceso actual</strong><span>Cómo registran equipos, alquileres y servicios, y dónde consultan esa información.</span></li><li><strong>Una demo concreta</strong><span>Elegir un equipo, consultar su cliente, abrir el historial y ver un reporte.</span></li><li><strong>Una primera etapa</strong><span>Definir qué proceso conviene abordar y qué tendría sentido desarrollar para su caso.</span></li></ul></section><section class="meeting"><h2>Conversemos 15 minutos.</h2><p>Reunión virtual o presencial en Lima, previa coordinación.<br>Responde al correo de Julio o <a href="https://wa.me/51936755465">escríbenos por WhatsApp</a> para acordar un horario.</p></section><p class="scope">La demo ilustra una forma de organizar la información. El alcance, los plazos y el precio se definen después de conocer la operación; no se presupone que todas las empresas necesiten las mismas funciones.</p><footer><div><strong>Julio · Nexo</strong><br><a href="mailto:asesoria@nexoamedida.com">asesoria@nexoamedida.com</a><br><a href="https://wa.me/51936755465">+51 936 755 465</a></div><div><a href="https://nexoamedida.com/#demo">nexoamedida.com · Explorar la demo</a><br><span>Jirón Tarapacá 670<br>Magdalena del Mar, Lima</span></div></footer></main></body></html>''',encoding='utf8')

text=['NEXO · CORREOS PARA CONSEGUIR UNA CONVERSACIÓN','USO INTERNO. Reemplaza todos los campos entre corchetes antes de enviar. Los horarios deben estar realmente disponibles. No enviar el ZIP ni esta guía a prospectos.','']
for item in CONTENT['emails']:
 text += ['='*65,item['title'],item['when'],'',('ASUNTO: '+item['subject'] if item.get('subject') else 'CANAL: WhatsApp'),'',item['body'],'']
KIT.joinpath('Correos-Nexo.txt').write_text('\n'.join(text),encoding='utf8')

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
 doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor.from_string('247459')
header=sec.header.paragraphs[0];header.text='NEXO  /  GUÍA INTERNA DE PROSPECCIÓN';header.style=doc.styles['Caption']
footer=sec.footer.paragraphs[0];footer.text='Julio · Nexo | asesoria@nexoamedida.com | Uso interno'
doc.add_heading('De empresa registrada a reunión confirmada',0)
doc.add_paragraph('Guía editable para trabajar tus primeras 100 empresas de alquiler. Contiene correos, derivación desde ventas, secuencia de seguimiento y guion de reunión.')
doc.add_paragraph('Adjunto por correo: Nexo-Informacion-Conectada.pdf. Adjunto por WhatsApp: Nexo-WhatsApp-Informacion.jpg. Esta guía y el ZIP son para tu trabajo interno.')
for section in CONTENT['guide']:
 doc.add_heading(section['title'],1)
 for paragraph in section['paragraphs']:doc.add_paragraph(paragraph)
 for bullet in section['bullets']:doc.add_paragraph(bullet,style='List Bullet')
 if section.get('script'):
  doc.add_heading('Guion sugerido',2);doc.add_paragraph(section['script'])
doc.add_page_break();doc.add_heading('Correos listos para adaptar',0)
for item in CONTENT['emails']:
 doc.add_heading(item['title'],1);doc.add_paragraph(item['when']);doc.add_paragraph(('Asunto: '+item['subject']) if item.get('subject') else 'Canal: WhatsApp',style='Heading 2')
 for line in item['body'].split('\n'):doc.add_paragraph(line)
 doc.add_paragraph('____________________________________________________________')
for style in ['Heading 1','Heading 2']:doc.styles[style].paragraph_format.keep_with_next=True
doc.save(KIT/'Nexo-Guia-y-correos-editables.docx')

copycards=[]
for item in CONTENT['emails']:
 subject=(f'<label for="{item["id"]}-subject">Asunto</label><input readonly id="{item["id"]}-subject" value="{e(item["subject"],quote=True)}"><button data-copy="{item["id"]}-subject">Copiar asunto</button>' if item.get('subject') else '')
 label='Mensaje de WhatsApp' if item.get('channel')=='whatsapp' else 'Correo'
 button='Copiar WhatsApp' if item.get('channel')=='whatsapp' else 'Copiar correo'
 copycards.append(f'<details class="email-card" {"open" if item["id"] in ["responsable","whatsapp"] else ""}><summary>{e(item["title"])}</summary><p class="help">{e(item["when"])}</p>{subject}<label for="{item["id"]}-body">{label}</label><textarea readonly id="{item["id"]}-body" rows="13">{e(item["body"])}</textarea><button data-copy="{item["id"]}-body">{button}</button></details>')
guides=[]
for item in CONTENT['guide']:
 guides.append('<details class="guide-card"><summary>'+e(item['title'])+'</summary>'+''.join('<p>'+e(p)+'</p>' for p in item['paragraphs'])+'<ul>'+''.join('<li>'+e(b)+'</li>' for b in item['bullets'])+'</ul>'+('<blockquote>'+e(item['script'])+'</blockquote>' if item.get('script') else '')+'</details>')
KIT.joinpath('index.html').write_text('''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow,noarchive"><meta name="referrer" content="no-referrer"><title>Nexo | Kit comercial</title><link rel="icon" href="/favicon.ico"><link rel="icon" href="/favicon.svg" type="image/svg+xml"><link rel="apple-touch-icon" href="/apple-touch-icon.png"><link rel="stylesheet" href="kit.css"></head><body><header><a href="/"><img src="/brand/nexo-logo.svg" alt="Nexo" width="145" height="51"></a><a class="back" href="/SKYROTE13/">Volver al CRM</a></header><main><h1>Información dispersa.<br>Una conversación para conectarla.</h1><p class="intro">Dos mensajes para usar hoy y un adjunto distinto para cada canal. Cambia los campos entre corchetes y ofrece horarios que tengas disponibles.</p><section aria-labelledby="downloads"><h2 id="downloads">Qué enviar en cada canal</h2><div class="downloads"><article class="client"><span>PARA ENVIAR POR CORREO</span><h3>Una ficha visual de una página</h3><p>Presenta a Nexo y muestra el antes y después de una consulta sobre un contrato. Incluye un enlace para coordinar una demostración de 15 minutos.</p><a class="button" download href="/presentacion/Nexo-Informacion-Conectada.pdf?v=20261020-2">Descargar PDF para correo</a><a class="preview" href="/presentacion/nexo-informacion-conectada.html?v=20261020-2" target="_blank" rel="noopener">Ver ficha en el navegador ↗</a></article><article><span>PARA ENVIAR POR WHATSAPP</span><h3>Una imagen directa para móvil</h3><p>Beneficio, presentación y una comparación clara en un JPG vertical de 1080 × 1350 px. Invita a responder al mensaje para coordinar la demostración.</p><a class="button" download href="/presentacion/Nexo-WhatsApp-Informacion.jpg?v=20261020-2">Descargar imagen para WhatsApp</a><a class="preview" href="/presentacion/Nexo-WhatsApp-Informacion.jpg?v=20261020-2" target="_blank" rel="noopener">Ver imagen ↗</a></article><article class="internal"><span>PARA TU USO INTERNO</span><h3>Guía y archivos editables</h3><p>Los dos mensajes de hoy, derivación desde ventas, seguimientos y guion de reunión. El ZIP reúne la imagen, el PDF y los documentos internos.</p><a class="button secondary" download href="Nexo-Guia-y-correos-editables.docx?v=20261020-2">Descargar guía Word</a><div class="secondary-links"><a download href="Correos-Nexo.txt">Mensajes en TXT</a><a download href="Nexo-Kit-Comercial.zip?v=20261020-2">Descargar todo en ZIP</a></div></article></div></section><section class="sequence"><h2>Si contactas hoy por ambos canales</h2><ol><li><strong>Correo:</strong> mensaje de hoy + PDF. Propón dos horarios reales.</li><li><strong>WhatsApp:</strong> mensaje breve + imagen. Si ya enviaste el correo, puedes mencionarlo.</li><li><strong>Si responden:</strong> continúa en ese canal y detén los seguimientos pendientes del otro.</li></ol><p>Cuando el contacto sea de ventas o informes, pide que te ayude a llegar a la persona responsable. Una reunión se confirma al acordar persona, fecha, hora y modalidad.</p></section><section aria-labelledby="emails"><h2 id="emails">Mensajes para copiar y adaptar</h2><p>Presentación, beneficio y una invitación concreta a ver el proyecto. <a href="https://github.com/julioroblesb/Demo-Sistema/blob/main/docs/comercial/investigacion-copy.md" target="_blank" rel="noopener">Leer la investigación y sus fuentes ↗</a> · <a download href="Investigacion-copy-Nexo.md?v=20261020-2">Descargar el análisis</a>.</p>'''+''.join(copycards)+'''</section><section aria-labelledby="guide"><h2 id="guide">Cómo conducir el contacto y la reunión</h2>'''+''.join(guides)+'''</section><p id="copy-status" role="status" aria-live="polite"></p></main><footer>Material interno de Nexo · Comparte con prospectos únicamente el PDF y la imagen preparados para cada canal.</footer><script src="kit.js"></script></body></html>''',encoding='utf8')
if not args.text_only:
 subprocess.run(['node','scripts/sales-kit/render.mjs'],cwd=ROOT,check=True)
 subprocess.run(['node','scripts/sales-kit/campaign.mjs'],cwd=ROOT,check=True)
with zipfile.ZipFile(KIT/'Nexo-Kit-Comercial.zip','w',zipfile.ZIP_DEFLATED) as archive:
 archive.write(PRESENT/'Nexo-Informacion-Conectada.pdf','Para-enviar/Correo-Nexo-Informacion-Conectada.pdf')
 archive.write(PRESENT/'Nexo-WhatsApp-Informacion.jpg','Para-enviar/WhatsApp-Nexo-Informacion.jpg')
 archive.write(KIT/'Nexo-Guia-y-correos-editables.docx','Uso-interno/Nexo-Guia-y-correos-editables.docx')
 archive.write(KIT/'Correos-Nexo.txt','Uso-interno/Correos-Nexo.txt')
 archive.write(KIT/'Investigacion-copy-Nexo.md','Uso-interno/Investigacion-copy-Nexo.md')
 archive.writestr('LEEME.txt','Nexo: envía el PDF por correo y la imagen JPG por WhatsApp, cada uno con el mensaje correspondiente. La guía Word, los mensajes TXT y el análisis de fuentes son para tu trabajo interno. Reemplaza los corchetes, verifica los datos de la empresa y ofrece horarios que tengas disponibles. El kit no envía correos ni crea reuniones automáticamente.')
print('Actualizados: guía Word, mensajes TXT, investigación, ZIP y página de consulta.' if args.text_only else 'Generados: presentación PDF, imagen JPG, guía Word, mensajes TXT, investigación, ZIP y página de consulta.')
