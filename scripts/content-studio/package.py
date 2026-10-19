from pathlib import Path
import json,zipfile,subprocess
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
root=Path(__file__).resolve().parents[2];out=root/'public/estudio-6e94b8f2c731';files=out/'archivos';p=json.loads((out/'plan.json').read_text())
for old in files.glob('plan-previo-*'):old.unlink()
wb=Workbook();ws=wb.active;ws.title='Calendario y copies'
ws.append(['ID','Fecha','Hora Lima','Formato','Título','Uso','Lanzamiento','Fijar','Copy Facebook','Copy Instagram','CTA','WhatsApp','Texto alternativo'])
for x in sorted(p['posts'],key=lambda x:(x['date'],x['time'])):ws.append([x['id'],x['date'],x['time'],x['format'],x['title'],x['adStatus'],'Sí' if x['launch'] else 'No','Sí' if x['pinned'] else 'No',x['fb'],x['ig'],x['cta'],x['whatsapp'],x['alt']])
ws2=wb.create_sheet('Historias');ws2.append(['ID','Fecha','Hora Lima','Tema','Contenido base','Instagram','Facebook'])
for s in p['storyItems']:ws2.append([s['id'],s['date'],s['time'],s['title'],s['content'],s['ig'],s['fb']])
ws3=wb.create_sheet('Pauta');ws3.append(['ID','Pieza','Uso','Formato','Ángulo','Copy','Titular','Palabra','Destino','Medición'])
for row in p['ads']:ws3.append(row)
ws4=wb.create_sheet('Instrucciones');ws4.append(['Tema','Instrucción'])
for a,b in [['Calendario vigente','20 de octubre al 13 de noviembre de 2026. Horas de Lima. Fechas sugeridas, sin publicación automática.'],['Lanzamiento','20 oct: P03 10:00, P02 13:00, P01 18:30. 21 oct: P05 12:30, P09 18:30. 22 oct: P06 18:30.'],['Fijar','P01, P02 y P03 donde la plataforma lo permita.'],['Pauta','Probar P02/A1 y P04/A2 primero; P08/A3 reserva. No se activó publicidad.'],['Video','MP4 9:16 y 4:5, 24 segundos, sin audio. Explicación editorial con datos ficticios.'],['Acceso','Portal público por enlace, sin contraseña. No guardar información confidencial.']]:ws4.append([a,b])
for sh in wb:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.fill=PatternFill('solid',fgColor='247459');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='top')
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top');cell.font=Font(size=11,color='202730')
  sh.row_dimensions[row[0].row].height=140 if sh==ws else 90
 for col in sh.columns:
  letter=col[0].column_letter;sh.column_dimensions[letter].width=50 if col[0].value in ['Copy Facebook','Copy Instagram','Contenido base','Copy','Instrucción'] else 27
wb.save(files/'Nexo-Grilla-Calendario-y-Copies.xlsx')
# Ensure media actually exists before offering downloadable packs.
for item in p['posts']+p['storyItems']:
 for asset in item['assets']:
  f=out/asset['file']
  if not f.is_file() or f.stat().st_size<1000:raise RuntimeError('Missing asset '+str(f))
def archive(name,paths):
 with zipfile.ZipFile(files/name,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
  for f in sorted(set(paths)):z.write(f,f.name)
 forcheck=zipfile.ZipFile(files/name).testzip()
 if forcheck:raise RuntimeError(forcheck)
for item in p['posts']:
 paths=[out/a['file'] for a in item['assets']]+[out/item['copyFile']]
 archive(item['id']+'-paquete.zip',paths)
for item in p['storyItems']:archive(item['id']+'-historias.zip',[out/a['file'] for a in item['assets']])
launch=[]
for item in p['posts']:
 if item['launch']:launch.extend([out/a['file'] for a in item['assets']]+[out/item['copyFile']])
launch.extend([files/'Guia-publicacion-Nexo.txt',files/'Calendario-Nexo.csv'])
archive('Nexo-Lanzamiento-6-publicaciones.zip',launch)
archive('Nexo-Contenido-Completo.zip',[f for f in files.iterdir() if f.suffix!='.zip'])
print('Packaged all assets; total MB:',round(sum(f.stat().st_size for f in files.iterdir())/1024**2,1))
